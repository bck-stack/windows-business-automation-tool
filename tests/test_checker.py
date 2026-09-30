import threading

import httpx
import openpyxl

import checker
from checker import CheckOptions, check_many, check_url, parse_url_list


def mock_client(handler):
    return httpx.Client(transport=httpx.MockTransport(handler), follow_redirects=True, max_redirects=3)


def handler(request: httpx.Request) -> httpx.Response:
    path = request.url.path
    if path == "/redirect":
        return httpx.Response(301, headers={"Location": "https://site.test/final"})
    if path == "/loop":
        return httpx.Response(302, headers={"Location": "https://site.test/loop"})
    if path == "/missing":
        return httpx.Response(404)
    if path == "/boom":
        return httpx.Response(503)
    if path == "/slow":
        raise httpx.ReadTimeout("slow", request=request)
    if request.url.host == "nodns.test":
        raise httpx.ConnectError("[Errno 11001] getaddrinfo failed", request=request)
    return httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"})


def test_parse_url_list():
    urls, rejected = parse_url_list("example.com\n# comment\nhttps://a.io, https://b.io\nHTTPS://A.IO\nnot a url\n\nfoo")
    assert urls == ["https://example.com", "https://a.io", "https://b.io"]
    assert rejected == ["not", "a", "url", "foo"]


def test_statuses_and_categories():
    opts = CheckOptions(retries=1)
    with mock_client(handler) as c:
        ok = check_url("https://site.test/", c, opts)
        red = check_url("https://site.test/redirect", c, opts)
        miss = check_url("https://site.test/missing", c, opts)
        boom = check_url("https://site.test/boom", c, opts)
        slow = check_url("https://site.test/slow", c, opts)
        loop = check_url("https://site.test/loop", c, opts)
        dns = check_url("https://nodns.test/", c, opts)
    assert (ok.status, ok.category, ok.content_type) == (200, "ok", "text/html")
    assert (red.status, red.category, red.redirects, red.final_url) == (200, "redirect", 1, "https://site.test/final")
    assert miss.category == "client" and boom.category == "server"
    assert slow.status == "TIMEOUT" and slow.category == "timeout"
    assert loop.status == "REDIRECT_LOOP"
    assert dns.status == "DNS_ERR"


def test_retries_transient_errors(monkeypatch):
    monkeypatch.setattr(checker.time, "sleep", lambda s: None)
    calls = {"n": 0}

    def flaky(request):
        calls["n"] += 1
        if calls["n"] == 1:
            raise httpx.ConnectError("reset", request=request)
        return httpx.Response(200)

    with mock_client(flaky) as c:
        assert check_url("https://x.test", c, CheckOptions(retries=1)).status == 200
    assert calls["n"] == 2


def test_check_many_parallel_order_and_cancel():
    seen = []
    with mock_client(handler) as c:
        urls = [f"https://site.test/{i}" for i in range(20)]
        results = check_many(urls, CheckOptions(workers=4), on_result=lambda i, r: seen.append(i), client=c)
        assert [r.url for r in results] == urls and len(seen) == 20

        cancel = threading.Event()
        cancel.set()
        assert check_many(urls, CheckOptions(workers=2), cancel=cancel, client=c) == []


def test_exports(tmp_path):
    with mock_client(handler) as c:
        results = check_many(["https://site.test/", "https://site.test/missing"], client=c)
    checker.export_csv(results, tmp_path / "r.csv")
    text = (tmp_path / "r.csv").read_text(encoding="utf-8-sig")
    assert text.splitlines()[0].startswith("URL,Status,Result")
    checker.export_xlsx(results, tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(tmp_path / "r.xlsx")
    assert wb["Results"].max_row == 3
    assert dict(wb["Summary"].iter_rows(min_row=2, values_only=True))["client"] == 1


def test_read_url_file_formats(tmp_path):
    (tmp_path / "u.csv").write_text("name,url\nA,https://a.io\nB,b.io\n", encoding="utf-8")
    urls, _ = parse_url_list(checker.read_url_file(tmp_path / "u.csv"))
    assert urls == ["https://a.io", "https://b.io"]
    wb = openpyxl.Workbook()
    wb.active.append(["https://x.io", None, "y.io"])
    wb.save(tmp_path / "u.xlsx")
    urls, _ = parse_url_list(checker.read_url_file(tmp_path / "u.xlsx"))
    assert urls == ["https://x.io", "https://y.io"]
